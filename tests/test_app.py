import io
import os
import unittest
from unittest.mock import patch, Mock

from openpyxl import load_workbook
from app import app


class WorkflowTests(unittest.TestCase):
    def test_api_errors_are_readable_json(self):
        client = app.test_client()
        for response, status in [
            (client.post('/api/analyze', data='{', content_type='application/json'), 400),
            (client.post('/api/import', data={'file': (io.BytesIO(b'title,price\n' + b'x' * 131073 + b',1000'), 'large-field.csv')}), 400),
            (client.post('/api/import', data={'file': (io.BytesIO(b'x' * (2 * 1024 * 1024)), 'oversized.csv')}), 413),
        ]:
            with self.subTest(status=status):
                self.assertEqual(response.status_code, status)
                self.assertTrue(response.json['error'])

    def setUp(self):
        self.client = app.test_client()
        self.payload = {'rows': [
            {'title': '동일 상품', 'source': 'A', 'price': 10000, 'shipping': 1000, 'group': 'A-50ml'},
            {'title': '동일 상품', 'source': 'B', 'price': 10500, 'shipping': 0, 'group': 'A-50ml'}],
            'settings': {'sale_price': 20, 'exchange_rate': 1000, 'fee_pct': 10,
                         'tax_pct': 5, 'refund_pct': 2, 'outbound_shipping': 3000, 'packaging': 500}}

    def test_landed_cost_and_profit(self):
        r = self.client.post('/api/analyze', json=self.payload)
        self.assertEqual(r.status_code, 200)
        cheap, expensive = r.json['rows']
        self.assertEqual(cheap['source'], 'B')
        self.assertEqual(cheap['rank'], 1)
        self.assertEqual(expensive['rank'], 2)
        self.assertEqual(cheap['profit'], 2600)
        self.assertEqual(cheap['total_cost'], 17400)
        self.assertAlmostEqual(cheap['roi_pct'], 2600 / 14000 * 100)
        self.assertEqual(cheap['margin_pct'], 13)
        self.assertAlmostEqual(cheap['break_even'], 14000 / .83)

    def test_csv_to_excel(self):
        r = self.client.post('/api/import', data={'file': (io.BytesIO(
            'title,price,source,shipping,group,url\n=1+1,1000,다이소,0,g,https://example.com\n'.encode()), 'data.csv')})
        self.assertEqual(r.status_code, 200)
        exported = self.client.post('/api/export', json={'rows': r.json['rows']})
        self.assertEqual(exported.status_code, 200)
        wb = load_workbook(io.BytesIO(exported.data))
        self.assertEqual(wb.active['C2'].value, "'=1+1")
        self.assertNotEqual(wb.active['C2'].data_type, 'f')
        self.assertEqual(wb.active['E2'].value, 1000)
        self.assertIn('비용 설정 및 안내', wb.sheetnames)

    def test_validation_and_unmatchable_rows(self):
        for price in [-1, 'NaN', 'abc', float('inf')]:
            r = self.client.post('/api/analyze', json={'rows': [{'title': 'x', 'price': price}]})
            self.assertEqual(r.status_code, 400)
        r = self.client.post('/api/analyze', json={'rows': [{'title': 'x', 'price': 0, 'url': 'javascript:alert(1)'}],
                                                 'settings': {'outbound_shipping': 0, 'packaging': 0, 'sale_price': 0}})
        row = r.json['rows'][0]
        self.assertIsNone(row['roi_pct'])
        self.assertIsNone(row['rank'])
        self.assertEqual(row['url'], '')

    @patch.dict(os.environ, {'NAVER_CLIENT_ID': '', 'NAVER_CLIENT_SECRET': ''})
    def test_search_without_credentials_is_explicit(self):
        r = self.client.get('/api/search?q=선크림')
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.json['rows'], [])
        self.assertEqual(len(r.json['links']), 8)

    @patch.dict(os.environ, {'NAVER_CLIENT_ID': 'test', 'NAVER_CLIENT_SECRET': 'test'})
    @patch('app.requests.get')
    def test_naver_response_mapping(self, get):
        get.return_value = Mock(json=lambda: {'items': [{'title': '<b>상품</b> &amp; 옵션', 'lprice': '1000',
                                                        'mallName': '몰', 'link': 'https://example.com'}]})
        r = self.client.get('/api/search?q=상품')
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.json['rows'][0]['title'], '상품 & 옵션')
        self.assertEqual(get.call_args.kwargs['params']['sort'], 'asc')

    def offer(self, group, market='PH_SHOPEE', **overrides):
        return dict(group=group, market=market, **{
            **dict(sale_price=1000, exchange_rate=20, fee_pct=10, tax_pct=5,
                   outbound_shipping=3000, other_cost=0, refund_pct=2, confirmed=True), **overrides})

    def test_market_currency_and_distinct_recommendations(self):
        data = dict(rows=[dict(title=g, group=g, source='다이소', price=p)
                          for g, p in [('A', 10000), ('A', 11000), ('B', 12000), ('C', 13000), ('D', 14000)]],
                    offers=[self.offer(g) for g in ['A', 'B', 'C', 'D']] + [
                        self.offer('A', 'VN_LAZADA', sale_price=600000, exchange_rate=.05)],
                    settings={'min_roi': 0})
        r = self.client.post('/api/compare', json=data)
        self.assertEqual(r.status_code, 200)
        self.assertEqual([x['group'] for x in r.json['top']], ['A', 'B', 'C'])
        self.assertEqual(r.json['top'][0]['currency'], 'VND')
        self.assertEqual(r.json['top'][0]['revenue'], 30000)
        self.assertEqual(r.json['top'][0]['profit'], 11400)
        exported = self.client.post('/api/compare/export', json=data)
        self.assertEqual(exported.status_code, 200)
        book = load_workbook(io.BytesIO(exported.data))
        self.assertEqual(book['추천 TOP3'].max_row, 4)
        self.assertEqual(book['추천 TOP3']['Q2'].value, 11400)
        self.assertIn('계산 가정', book.sheetnames)

    def test_unconfirmed_unmatched_and_missing_costs(self):
        data = dict(rows=[dict(title='A', group='A', price=1000)], offers=[
            self.offer('A', confirmed=False), self.offer('missing')])
        r = self.client.post('/api/compare', json=data)
        self.assertEqual(r.json['top'], [])
        self.assertEqual(r.json['unmatched_groups'], ['missing'])
        del data['offers'][0]['fee_pct']
        self.assertEqual(self.client.post('/api/compare', json=data).status_code, 400)

    def test_profit_and_roi_rank_differ_and_exclude_losses(self):
        data = dict(rows=[dict(title='A', group='A', price=10000), dict(title='B', group='B', price=1000),
                          dict(title='C', group='C', price=100000)], settings={'min_roi': 0},
                    offers=[self.offer('A', sale_price=1600), self.offer('B', sale_price=1000), self.offer('C')])
        self.assertEqual(self.client.post('/api/compare', json=data).json['top'][0]['group'], 'A')
        data['metric'] = 'roi_pct'
        result = self.client.post('/api/compare', json=data).json
        self.assertEqual(result['top'][0]['group'], 'B')
        self.assertNotIn('C', [r['group'] for r in result['top']])

    def test_home_renders(self):
        response = self.client.get('/')
        self.assertEqual(response.status_code, 200)
        self.assertIn('사입 레이더', response.get_data(as_text=True))

    def test_json_object_and_settings_validation(self):
        endpoints = ('/api/analyze', '/api/compare', '/api/export', '/api/compare/export')
        for endpoint in endpoints:
            for body in ('null', '[]', '[1]', '"상품"'):
                with self.subTest(endpoint=endpoint, body=body):
                    response = self.client.post(endpoint, data=body, content_type='application/json')
                    self.assertEqual(response.status_code, 400)
                    self.assertTrue(response.is_json)
                    self.assertIn('error', response.json)
            for invalid_settings in (None, [], [1]):
                with self.subTest(endpoint=endpoint, settings=invalid_settings):
                    response = self.client.post(endpoint, json={'settings': invalid_settings})
                    self.assertEqual(response.status_code, 400)
                    self.assertTrue(response.is_json)
                    self.assertIn('error', response.json)

    def test_market_must_be_a_supported_string(self):
        for endpoint in ('/api/compare', '/api/compare/export'):
            with self.subTest(endpoint=endpoint):
                response = self.client.post(endpoint, json={
                    **self.payload, 'offers': [self.offer('A-50ml', market=['PH_SHOPEE'])]})
                self.assertEqual(response.status_code, 400)
                self.assertTrue(response.is_json)
                self.assertIn('error', response.json)

    def test_excel_export_cleans_unsupported_control_characters(self):
        payload = {**self.payload, 'rows': [{**self.payload['rows'][0], 'title': '상품\x00명\x0b'}],
                   'offers': [self.offer('A-50ml')]}
        for endpoint, sheet in (('/api/export', '상품 수익 분석'),
                                ('/api/compare/export', '국가 플랫폼 비교')):
            with self.subTest(endpoint=endpoint):
                response = self.client.post(endpoint, json=payload)
                self.assertEqual(response.status_code, 200)
                workbook = load_workbook(io.BytesIO(response.data))
                self.assertEqual(workbook[sheet]['C2'].value, '상품명')

    @patch.dict(os.environ, {'NAVER_CLIENT_ID': 'test', 'NAVER_CLIENT_SECRET': 'test'})
    @patch('app.requests.get')
    def test_malformed_naver_response_returns_json_gateway_error(self, get):
        responses = [
            {'items': None},
            {'items': [{'title': '상품', 'lprice': '1000'}]},
            {'items': [{'title': '상품', 'lprice': '숫자 아님', 'link': 'https://example.com'}]},
        ]
        for body in responses:
            with self.subTest(body=body):
                get.return_value = Mock(json=lambda value=body: value)
                response = self.client.get('/api/search?q=상품')
                self.assertEqual(response.status_code, 502)
                self.assertTrue(response.is_json)
                self.assertIn('error', response.json)


if __name__ == '__main__':
    unittest.main()
