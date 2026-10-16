"""Validate the actual browser-generated XLSX with an independent Excel reader.

Run `node tests/test_web_xlsx.cjs` first, then `.venv/bin/python tests/check_web_xlsx.py`.
"""

from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree

from openpyxl import load_workbook


path = Path("/tmp/web-xlsx-test.xlsx")
with ZipFile(path) as archive:
    assert archive.testzip() is None, "ZIP CRC verification failed"
    for filename in archive.namelist():
        if filename.endswith((".xml", ".rels")):
            ElementTree.fromstring(archive.read(filename))

workbook = load_workbook(path, data_only=False)
assert workbook.sheetnames == ["추천 상품 🇰🇷", "비교   수익", "비교   수익 (2)"]
sheet = workbook.worksheets[0]
assert sheet["A2"].value == "라운드랩 <독도> & 크림"
assert sheet["A3"].value == "베트남 Tiếng Việt / 필리핀"
assert sheet["A4"].value == "  공백\t유지\n줄바꿈  "
assert sheet["A5"].value == "줄바꿈\r\n원문 유지"
assert sheet["B2"].value == 1234.567
assert sheet["B2"].data_type == "n"
assert sheet["B3"].value == -42.5
assert sheet["B4"].value == 0
assert sheet["B5"].value == 1e-7
assert sheet["C2"].value is None
assert sheet["D2"].value is True
assert sheet["D3"].value is False
assert sheet["D2"].data_type == "b"
assert sheet["E2"].value == '=HYPERLINK("https://example.com","클릭")'
assert sheet["E2"].data_type == "s", "Formula-shaped input must remain literal text"
assert sheet["E3"].value == "+SUM(A1:A2)"
assert sheet["E3"].data_type == "s"
assert sheet["E4"].value == "위험문자제거"
assert sheet.freeze_panes == "A2"
assert sheet["A1"].font.bold is True
assert workbook.worksheets[1]["A2"].value == 9900
assert workbook.worksheets[2].max_row == 1
assert all(cell.data_type != "f" for row in sheet for cell in row)
workbook.close()
print("openpyxl validated Unicode, numeric precision, blanks, booleans, formula safety, and workbook structure.")
