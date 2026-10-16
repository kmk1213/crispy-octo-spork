import csv
import io
import math
import os
from urllib.parse import urlencode, urlsplit

import requests
from flask import Flask, jsonify, render_template, request, send_file
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from werkzeug.exceptions import HTTPException

app = Flask(__name__)
app.config['MAX_CONTENT_LENGTH'] = 2 * 1024 * 1024
PLATFORMS = {
    '올리브영': 'https://www.oliveyoung.co.kr/store/search/getSearchMain.do?query=',
    '다이소': 'https://www.daisomall.co.kr/ds/search/totalSearch?searchTerm=',
    '쿠팡': 'https://www.coupang.com/np/search?q=',
    '네이버쇼핑': 'https://search.shopping.naver.com/search/all?query=',
}
MARKETS = {
    'PH_SHOPEE': dict(country='필리핀', platform='Shopee', currency='PHP', search='https://shopee.ph/search?keyword='),
    'PH_LAZADA': dict(country='필리핀', platform='Lazada', currency='PHP', search='https://www.lazada.com.ph/catalog/?q='),
    'VN_SHOPEE': dict(country='베트남', platform='Shopee', currency='VND', search='https://shopee.vn/search?keyword='),
    'VN_LAZADA': dict(country='베트남', platform='Lazada', currency='VND', search='https://www.lazada.vn/catalog/?q='),
}
DEFAULTS = dict(sale_price=20000, exchange_rate=1, fee_pct=12, tax_pct=0,
                purchase_shipping=0, outbound_shipping=3000, packaging=500,
                other_cost=0, refund_pct=0, min_roi=10)


def number(value, name):
    try:
        n = float(value)
    except (TypeError, ValueError):
        raise ValueError(f'{name}: 숫자를 입력하세요.')
    if not math.isfinite(n) or n < 0 or n > 1e12:
        raise ValueError(f'{name}: 유효한 0 이상의 숫자를 입력하세요.')
    return n


def settings(data):
    if not isinstance(data, dict):
        raise ValueError('비용 설정은 항목과 숫자로 된 객체여야 합니다.')
    result = {k: number(data.get(k, v), k) for k, v in DEFAULTS.items()}
    for k in ('fee_pct', 'tax_pct', 'refund_pct'):
        if result[k] > 100:
            raise ValueError(f'{k}: 100% 이하여야 합니다.')
    if result['exchange_rate'] == 0:
        raise ValueError('환율은 0보다 커야 합니다.')
    return result


def calculate(row, s):
    price = number(row.get('price'), '사입가')
    revenue = s['sale_price'] * s['exchange_rate']
    shipping = number(row.get('shipping') or 0, '판매처 배송료')
    purchase = price + shipping + s['purchase_shipping']
    fixed = purchase + s['outbound_shipping'] + s['packaging'] + s['other_cost']
    rate = (s['fee_pct'] + s['tax_pct'] + s['refund_pct']) / 100
    fee, tax, reserve = (revenue * s[k] / 100 for k in ('fee_pct', 'tax_pct', 'refund_pct'))
    profit = revenue - fixed - fee - tax - reserve
    cost = fixed + fee + tax + reserve
    return {**row, 'price': price, 'shipping': shipping, 'purchase_total': purchase,
            'revenue': revenue, 'fee': fee, 'tax': tax, 'refund_reserve': reserve,
            'total_cost': cost, 'profit': profit,
            'roi_pct': profit / fixed * 100 if fixed else None,
            'margin_pct': profit / revenue * 100 if revenue else None,
            'break_even': fixed / (1 - rate) if rate < 1 else None}


def safe_url(value):
    value = str(value or '')
    return value if urlsplit(value).scheme in ('http', 'https') else ''


def analyze(data):
    if not isinstance(data, dict):
        raise ValueError('요청 데이터는 상품 목록과 비용 설정을 담은 객체여야 합니다.')
    s = settings(data.get('settings', {}))
    rows = data.get('rows', [])
    if not isinstance(rows, list) or len(rows) > 1000:
        raise ValueError('상품은 최대 1,000개입니다.')
    output = []
    for row in rows:
        if not isinstance(row, dict) or not str(row.get('title', '')).strip():
            raise ValueError('상품명이 필요합니다.')
        clean = {k: str(row.get(k, ''))[:2000] for k in ('title', 'source', 'group')}
        clean.update(price=row.get('price'), shipping=row.get('shipping'), url=safe_url(row.get('url')))
        output.append(calculate(clean, s))
    output.sort(key=lambda r: r['purchase_total'])
    groups = {}
    for r in output:
        group = r['group'].strip()
        r['rank'] = None
        if group:
            groups[group] = groups.get(group, 0) + 1
            r['rank'] = groups[group]
        r['recommended'] = r['profit'] > 0 and r['roi_pct'] is not None and r['roi_pct'] >= s['min_roi']
    return output, s


@app.errorhandler(ValueError)
def invalid(error):
    return jsonify(error=str(error)), 400


@app.errorhandler(HTTPException)
def http_error(error):
    if not request.path.startswith('/api/'):
        return error
    messages = {400: '요청 형식을 확인하세요. JSON 데이터가 올바르지 않습니다.',
                413: '파일이 너무 큽니다. 2MB 이하 CSV를 사용하세요.',
                415: 'JSON 형식으로 요청하세요.', 404: '요청한 기능을 찾을 수 없습니다.',
                405: '지원하지 않는 요청 방식입니다.'}
    return jsonify(error=messages.get(error.code, '요청을 처리하지 못했습니다.')), error.code


@app.errorhandler(500)
def server_error(error):
    # Flask logs the original traceback. Do not expose request contents or credentials.
    if request.path.startswith('/api/'):
        return jsonify(error='서버 처리 중 오류가 발생했습니다. 잠시 후 다시 시도하세요.'), 500
    return '<html lang="ko"><meta charset="utf-8"><h1>화면을 불러오지 못했습니다.</h1><p>잠시 후 새로고침해 주세요.</p></html>', 500


@app.get('/')
def home():
    return render_template('index.html', defaults=DEFAULTS, markets=MARKETS)


@app.get('/api/search')
def search():
    keyword = request.args.get('q', '').strip()[:100]
    if not keyword:
        raise ValueError('검색어를 입력하세요.')
    links = {name: base + urlencode({'q': keyword})[2:] for name, base in PLATFORMS.items()}
    links.update({m['country'] + ' ' + m['platform']: m['search'] + urlencode({'q': keyword})[2:]
                  for m in MARKETS.values()})
    client, secret = os.getenv('NAVER_CLIENT_ID'), os.getenv('NAVER_CLIENT_SECRET')
    if not client or not secret:
        return jsonify(rows=[], links=links, message='네이버 API 키가 없어 자동 가격 조회를 실행하지 않았습니다. 검색 링크 또는 CSV·직접 입력을 사용하세요.')
    try:
        response = requests.get('https://openapi.naver.com/v1/search/shop.json',
            params=dict(query=keyword, display=100, sort='asc'),
            headers={'X-Naver-Client-Id': client, 'X-Naver-Client-Secret': secret}, timeout=20)
        response.raise_for_status()
        body = response.json()
        if not isinstance(body, dict) or not isinstance(body.get('items'), list):
            raise ValueError('Invalid search response')
        import html
        import re
        rows = []
        for item in body['items']:
            if (not isinstance(item, dict) or not isinstance(item.get('title'), str)
                    or not isinstance(item.get('link'), str) or 'lprice' not in item):
                raise ValueError('Invalid search item')
            price = number(item['lprice'], '검색 가격')
            if price == 0:
                continue
            rows.append(dict(title=html.unescape(re.sub('<[^>]*>', '', item['title'])), price=price,
                             source='네이버쇼핑 / ' + str(item.get('mallName') or ''), shipping=0,
                             group='', url=safe_url(item.get('link'))))
    except (requests.RequestException, ValueError):
        return jsonify(error='네이버 검색 API 요청에 실패했습니다. 키·접근 권한·호출 한도를 확인하세요.'), 502
    return jsonify(rows=rows, links=links, message='네이버 표시 가격입니다. 배송료 미포함·옵션 및 동일 상품 여부 확인이 필요합니다. 다른 플랫폼 직접 수집은 지원하지 않습니다.')


@app.post('/api/import')
def import_csv():
    file = request.files.get('file')
    if not file:
        raise ValueError('CSV 파일이 필요합니다.')
    try:
        reader = csv.DictReader(io.StringIO(file.read().decode('utf-8-sig')))
        if not {'title', 'price'}.issubset(reader.fieldnames or []):
            raise ValueError('CSV에 title, price 열이 필요합니다.')
        rows = list(reader)
        analyze({'rows': rows})
    except UnicodeDecodeError:
        raise ValueError('UTF-8 CSV 파일을 사용하세요.')
    except csv.Error:
        raise ValueError('CSV 열 형식이나 항목 길이를 확인하세요. 상품별 항목은 128KB 이하여야 합니다.')
    return jsonify(rows=rows)


@app.post('/api/analyze')
def analysis():
    rows, _ = analyze(request.get_json())
    return jsonify(rows=rows)


def cell(value):
    if isinstance(value, str):
        value = ILLEGAL_CHARACTERS_RE.sub('', value)
    if isinstance(value, str) and value.startswith(('=', '+', '-', '@', '\t', '\r', '\n')):
        return "'" + value
    return value


def compare(data):
    rows, common = analyze(data)
    offers = data.get('offers', [])
    if not isinstance(offers, list) or len(offers) > 200:
        raise ValueError('판매 시나리오는 최대 200개입니다.')
    metric = data.get('metric', 'profit')
    if metric not in ('profit', 'roi_pct'):
        raise ValueError('추천 기준은 profit 또는 roi_pct입니다.')
    results = []
    unmatched = []
    for index, offer in enumerate(offers):
        if (not isinstance(offer, dict) or not isinstance(offer.get('market'), str)
                or offer['market'] not in MARKETS):
            raise ValueError('필리핀·베트남의 지원 판매 플랫폼을 선택하세요.')
        group = str(offer.get('group', '')).strip()
        if not group:
            raise ValueError('판매 시나리오에 동일상품 그룹이 필요합니다.')
        # Destination-specific prices and costs must be explicit: no invented market defaults.
        required = ('sale_price', 'exchange_rate', 'fee_pct', 'tax_pct', 'outbound_shipping', 'other_cost', 'refund_pct')
        for key in required:
            if key not in offer or str(offer[key]).strip() == '':
                raise ValueError(f'판매 시나리오 {index + 1}: {key} 값을 입력하세요. 해당 없으면 0입니다.')
        s = settings({**common, **{key: offer[key] for key in required}})
        matches = [row for row in rows if row['group'].strip() == group]
        if not matches:
            unmatched.append(group)
        for row in matches:
            result = calculate(row, s)
            market = MARKETS[offer['market']]
            result.update(market=offer['market'], country=market['country'], platform=market['platform'],
                          currency=market['currency'], sale_price=s['sale_price'], exchange_rate=s['exchange_rate'],
                          sale_url=safe_url(offer.get('sale_url')), confirmed=offer.get('confirmed') is True,
                          offer_index=index, assumptions=s, recommendation_rank=None)
            result['recommended'] = (result['confirmed'] and result['profit'] > 0 and
                result['roi_pct'] is not None and result['roi_pct'] >= s['min_roi'])
            results.append(result)
    # Recommend distinct products; multiple suppliers/markets must not occupy all three places.
    candidates = sorted((r for r in results if r['recommended']),
                        key=lambda r: (-r[metric], -r['profit'], r['purchase_total'], r['group']))
    top, seen = [], set()
    for row in candidates:
        if row['group'] in seen:
            continue
        seen.add(row['group'])
        row['recommendation_rank'] = len(top) + 1
        top.append(row)
        if len(top) == 3:
            break
    results.sort(key=lambda r: -r['profit'])
    return dict(rows=results, top=top, unmatched_groups=sorted(set(unmatched)), metric=metric,
                note='입력한 판매가·비용 기준 예상 수익 추천입니다. 실제 수요·판매량·실시간 시장 최저가는 검증하지 않았습니다.')


@app.post('/api/compare')
def comparison():
    return jsonify(compare(request.get_json()))


@app.post('/api/compare/export')
def export_comparison():
    data = compare(request.get_json())
    wb = Workbook()
    fields = [('recommendation_rank', '추천 상품 순위'), ('group', '동일상품 그룹'), ('title', '상품'),
              ('source', '사입처'), ('price', '사입가 KRW'), ('purchase_total', '사입 합계 KRW'),
              ('country', '판매 국가'), ('platform', '판매 플랫폼'), ('currency', '판매 통화'),
              ('sale_price', '판매가 현지통화'), ('exchange_rate', '환율 KRW/현지통화'),
              ('revenue', '매출 KRW'), ('fee', '수수료 KRW'), ('tax', '세금 충당 KRW'),
              ('refund_reserve', '반품 충당 KRW'), ('total_cost', '전체비용 KRW'),
              ('profit', '순이익 KRW'), ('roi_pct', '수익률 %'), ('margin_pct', '마진율 %'),
              ('confirmed', '동일상품·시세 사용자 확인'), ('url', '사입 링크'), ('sale_url', '판매 시세 링크')]
    for name, rows in [('추천 TOP3', data['top']), ('국가 플랫폼 비교', data['rows'])]:
        sheet = wb.create_sheet(name)
        sheet.append([label for _, label in fields])
        for row in rows:
            sheet.append([cell(row.get(key)) for key, _ in fields])
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = 24
    wb.remove(wb['Sheet'])
    sheet = wb.create_sheet('계산 가정')
    sheet.append(['시나리오 번호', '동일상품 그룹', '국가', '플랫폼', *DEFAULTS.keys()])
    seen = set()
    for row in data['rows']:
        if row['offer_index'] in seen:
            continue
        seen.add(row['offer_index'])
        sheet.append([row['offer_index'] + 1, cell(row['group']), row['country'], row['platform'],
                      *[row['assumptions'][key] for key in DEFAULTS]])
    sheet.append(['추천 기준', data['metric']])
    sheet.append([data['note']])
    sheet.append(['세금은 매출 비례 충당 추정. 실제 세무 정산 아님. 비용은 상품 1개당 원화.'])
    sheet.append(['수익률 = 순이익 / (사입 합계 + 고객 발송비 + 포장비 + 기타비) × 100.'])
    sheet.append(['그룹이 일치하는 후보에만 시나리오를 적용. 사용자 확인 및 목표 수익률 충족 상품만 TOP3에 포함.'])
    if data['unmatched_groups']:
        sheet.append(['매칭되지 않은 그룹', cell(', '.join(data['unmatched_groups']))])
    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    return send_file(stream, as_attachment=True, download_name='PH-VN-product-comparison.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@app.post('/api/export')
def export():
    rows, s = analyze(request.get_json())
    wb = Workbook()
    ws = wb.active
    ws.title = '상품 수익 분석'
    fields = [('rank','그룹 내 가격순위'), ('group','동일상품 그룹'), ('title','상품명'), ('source','판매처'),
              ('price','사입가 KRW'), ('shipping','판매처 배송료'), ('purchase_total','사입 합계'),
              ('revenue','판매 매출 KRW'), ('fee','수수료'), ('tax','세금 충당액'), ('refund_reserve','반품 충당액'),
              ('total_cost','전체 비용'), ('profit','예상 순이익'), ('roi_pct','투입비용 대비 수익률 %'),
              ('margin_pct','매출 대비 마진율 %'), ('break_even','손익분기 매출 KRW'), ('recommended','목표 충족'), ('url','상품 링크')]
    ws.append([label for _, label in fields])
    for row in rows:
        ws.append([cell(row.get(key)) for key, _ in fields])
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 22
    assumptions = wb.create_sheet('비용 설정 및 안내')
    assumptions.append(['항목', '값'])
    for key, val in s.items():
        assumptions.append([key, val])
    for note in ['모든 비용은 상품 1개 기준. 사입가와 비용은 원화, 판매가는 입력 통화 × 환율.',
                 '수수료·세금·반품 충당액은 원화 매출에 각 비율을 곱한 추정값.',
                 '수익률 = 순이익 / (사입 합계 + 발송비 + 포장비 + 기타비) × 100.',
                 '세금은 단순 충당 추정이며 실제 부가세·소득세·관세 정산이 아닙니다. 관세 등은 기타비에 반영하세요.',
                 '동일상품 그룹은 모델·용량·수량·옵션이 같은지 사용자가 확인한 후 입력합니다.',
                 '가격순위는 입력·조회한 후보 안에서만 계산하며 시장 전체 최저가를 보장하지 않습니다.']:
        assumptions.append([note])
    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    return send_file(stream, as_attachment=True, download_name='product-profit.xlsx', mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


if __name__ == '__main__':
    app.run(host='127.0.0.1', port=int(os.getenv('PORT', '8000')))
